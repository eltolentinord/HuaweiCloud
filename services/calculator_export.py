# coding: utf-8
"""Exportación de la lista de la calculadora: Excel, CSV y PDF (tamaño carta).

Recibe la lista YA recalculada en el servidor (``costs.calculator.price_list``): ningún
importe viene del navegador. Sin precio => "Precio no disponible", nunca 0.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

from exports.excel import neutralize_formulas
from services.cost_pdf import AZUL_MEDIO, ESTILOS, GRIS_TEXTO, LINEA, _kpis, _p, _tabla, _texto

TITULO = "Presupuesto estimado"
PROVEEDOR = "Huawei Cloud"
PRODUCTOS = {"ecs": "Servidor (ECS)", "evs": "Disco (EVS)", "eip": "IP pública (EIP)",
             "rds": "Base de datos (RDS)", "obs": "Almacenamiento (OBS)"}
MODOS = {"monthly": "Mensual", "yearly": "Anual", "on_demand": "Por uso"}
NA = "Precio no disponible"
NOTA = ("Precios oficiales de lista de Huawei Cloud (BSS, official_website_amount) consultados con la cuenta; "
        "no incluyen descuentos, impuestos ni consumos variables no indicados. Simulación: no crea recursos.")


def _region(regions: Dict[str, str], region: str) -> str:
    name = regions.get(region)
    return f"{name} ({region})" if name and name != region else region


def _rows(priced: Dict[str, Any], regions: Dict[str, str]) -> List[List[Any]]:
    rows = []
    for n, item in enumerate(priced["items"], start=1):
        rows.append([n, item.get("name") or PRODUCTOS[item["product"]], _region(regions, item["region"]),
                     item["summary"], MODOS[item["billing_mode"]], item["duration_text"], item["quantity"],
                     item["total"] if item["total"] is not None else NA, item["currency"] or ""])
    return rows


HEAD = ["#", "Producto", "Región", "Configuración", "Cobro", "Duración", "Cantidad", "Subtotal", "Moneda"]
DETAIL_HEAD = ["#", "Componente", "Precio unitario", "Unidad", "× Duración/uso", "× Cantidad", "Subtotal",
               "Moneda", "Fuente", "Motivo si no hay precio"]


def _detail(priced: Dict[str, Any]) -> List[List[Any]]:
    rows = []
    for n, item in enumerate(priced["items"], start=1):
        for line in item["lines"]:
            rows.append([n, line["label"], line["unit_amount"] or NA, line["unit"] or "",
                         line["multiplier"] if line["available"] else "",
                         item["quantity"] * line["units_per_item"], line["subtotal"] or NA,
                         line["currency"] or "", line["source"] or "", line["reason"] or ""])
    return rows


def export_xlsx(priced: Dict[str, Any], *, regions: Dict[str, str], generated_at: datetime) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"{TITULO} · {PROVEEDOR}"])
    ws.append([f"Generado: {generated_at:%Y-%m-%d %H:%M}"])
    ws.append([])
    ws.append(HEAD)
    header_row = ws.max_row
    for row in _rows(priced, regions):
        ws.append(row)
    ws.append([])
    for t in priced["totals"]:
        ws.append(["", "Total estimado", "", "", "", "", "", t["total"], t["currency"]])
    if priced["incomplete"]:
        ws.append(["", f"{priced['incomplete']} ítem(s) sin precio completo: no se suman al total."])
    ws.append([])
    ws.append([NOTA])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center")
    ws["A1"].font = Font(bold=True, size=14)
    for col, width in zip("ABCDEFGHI", (5, 22, 26, 60, 12, 12, 10, 16, 9)):
        ws.column_dimensions[col].width = width
    detail = wb.create_sheet("Detalle")
    detail.append(DETAIL_HEAD)
    for row in _detail(priced):
        detail.append(row)
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for col, width in zip("ABCDEFGHIJ", (5, 40, 16, 8, 14, 10, 14, 9, 50, 60)):
        detail.column_dimensions[col].width = width
    neutralize_formulas(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_csv(priced: Dict[str, Any], *, regions: Dict[str, str], generated_at: datetime) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(HEAD)
    for row in _rows(priced, regions):
        writer.writerow([_safe_cell(v) for v in row])
    for t in priced["totals"]:
        writer.writerow(["", "Total estimado", "", "", "", "", "", t["total"], t["currency"]])
    return ("﻿" + out.getvalue()).encode("utf-8")


def _safe_cell(value: Any) -> Any:
    """Evita fórmulas al abrir el CSV en Excel (mismo criterio que exports/excel)."""
    text = str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") and not text.replace(".", "", 1).lstrip("-").isdigit() else value


class _Pagina:
    def __init__(self, generado: str):
        self.generado = generado

    def __call__(self, canvas, doc):
        canvas.saveState()
        ancho, alto = doc.pagesize
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(AZUL_MEDIO)
        canvas.drawRightString(ancho - doc.rightMargin, alto - 12 * mm, f"{TITULO} · {PROVEEDOR}")
        canvas.setStrokeColor(LINEA)
        canvas.line(doc.leftMargin, alto - 13.5 * mm, ancho - doc.rightMargin, alto - 13.5 * mm)
        canvas.setFillColor(GRIS_TEXTO)
        canvas.drawString(doc.leftMargin, 10 * mm, f"Generado: {self.generado}")
        canvas.drawRightString(ancho - doc.rightMargin, 10 * mm, f"Página {doc.page}")
        canvas.restoreState()


def export_pdf(priced: Dict[str, Any], *, regions: Dict[str, str], generated_at: datetime) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=LETTER, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=22 * mm,
                            bottomMargin=20 * mm, title=f"{TITULO} - {PROVEEDOR}", author=PROVEEDOR)
    ancho = doc.width
    historia: List[Any] = [_p(f"{TITULO}", "h1"),
                           _p(f"{PROVEEDOR} · generado el {generated_at:%Y-%m-%d %H:%M}", "nota"), Spacer(1, 3 * mm)]
    kpis = [(f"Total ({t['currency']})", _texto(f"{t['currency']} {float(t['total']):,.2f}")) for t in priced["totals"]]
    kpis.append(("Ítems con precio", _texto(f"{priced['complete']} de {priced['complete'] + priced['incomplete']}")))
    historia.append(_kpis(kpis[:4], ancho))
    historia.append(_p("Resumen de la lista", "h2"))
    filas = [[_p(n), _p(item.get("name") or PRODUCTOS[item["product"]]), _p(_region(regions, item["region"])),
              _p(item["summary"]), _p(f"{MODOS[item['billing_mode']]} · {item['duration_text']}"), _p(item["quantity"], "celda_der"),
              _p(f"{item['currency']} {float(item['total']):,.2f}" if item["total"] is not None else NA, "celda_der")]
             for n, item in enumerate(priced["items"], start=1)]
    historia.append(_tabla(["#", "Producto", "Región", "Configuración", "Cobro", "Cant.", "Subtotal"], filas,
                           [ancho * 0.05, ancho * 0.14, ancho * 0.17, ancho * 0.30, ancho * 0.13, ancho * 0.07,
                            ancho * 0.14]))
    if priced["incomplete"]:
        historia.append(_p(f"{priced['incomplete']} ítem(s) sin precio completo no se suman al total.", "nota"))
    for n, item in enumerate(priced["items"], start=1):
        filas = [[_p(l["label"]), _p(f"{l['currency']} {l['unit_amount']}/{l['unit']}" if l["available"] else NA),
                  _p(f"× {l['multiplier']}" + (f" × {item['quantity'] * l['units_per_item']}"
                                              if item["quantity"] * l["units_per_item"] != 1 else "") if l["available"] else
                     (l["reason"] or "")),
                  _p(f"{l['currency']} {float(l['subtotal']):,.2f}" if l["subtotal"] else NA, "celda_der")]
                 for l in item["lines"]]
        titulo = f"{n}. {item.get('name') or PRODUCTOS[item['product']]} — {_region(regions, item['region'])}"
        historia.append(KeepTogether([_p(titulo, "h2"),
                                      _p(f"{MODOS[item['billing_mode']]} · {item['duration_text']} · "
                                         f"cantidad {item['quantity']}", "intro"),
                                      _tabla(["Componente", "Precio unitario", "Cálculo", "Subtotal"], filas,
                                             [ancho * 0.36, ancho * 0.2, ancho * 0.28, ancho * 0.16])]))
        for aviso in item["warnings"]:
            historia.append(_p(aviso, "nota"))
    historia += [Spacer(1, 4 * mm), Paragraph(_texto(NOTA), ESTILOS["nota"])]
    doc.build(historia, onFirstPage=_Pagina(f"{generated_at:%Y-%m-%d %H:%M}"),
              onLaterPages=_Pagina(f"{generated_at:%Y-%m-%d %H:%M}"))
    return out.getvalue()
