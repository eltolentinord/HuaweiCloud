# coding: utf-8
"""Reportes de auditoría de servidores: comparativo HTML (misma lógica que el
``compare-audits.sh`` del usuario), Excel y PDF (tamaño carta, estilo de la plataforma).

Entrada: la lista de resultados ``{server_name, data}`` donde ``data`` es el JSON de
``server-audit.sh -j``. Nada se recalcula ni se inventa: puntajes y controles vienen del script.
"""

from __future__ import annotations

import html
import io
from datetime import datetime
from typing import Any, Dict, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, SimpleDocTemplate, Spacer

from exports.excel import neutralize_formulas
from services.cost_pdf import AZUL_MEDIO, GRIS_TEXTO, LINEA, _kpis, _p, _tabla, _texto, _vinetas

TITULO = "Auditoría de servidores"
STATUS_ORDER = {"FAIL": 0, "WARN": 1, "INFO": 2, "OK": 3}


def _sorted(results: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return sorted(results, key=lambda r: -(r["data"].get("overall_score") or 0))


def _checks(data: Dict[str, Any]) -> List[Dict[str, Any]]:
    return [c for c in data.get("checks") or [] if isinstance(c, dict)]


# ---------------------------------------------------------------- HTML (compare-audits.sh)
def comparative_html(results: Sequence[Dict[str, Any]], *, generated_at: datetime) -> str:
    """Misma lógica que ``compare-audits.sh`` (ordenado por puntaje, resumen y control por control),
    en Python porque la plataforma corre en Windows sin bash."""
    return _python_comparative(results, generated_at=generated_at)


def _bar_color(pct: int) -> str:
    return "#16a34a" if pct >= 80 else "#f59e0b" if pct >= 50 else "#dc2626"


def _python_comparative(results: Sequence[Dict[str, Any]], *, generated_at: datetime) -> str:
    esc = lambda v: html.escape(str(v), quote=True)  # noqa: E731
    servers = _sorted(results)
    now = generated_at.strftime("%Y-%m-%d %H:%M:%S")

    def bar(pct: Any, label: str) -> str:
        pct = int(pct or 0)
        return (f'<div class="bar-row"><div class="bar-track"><div class="bar-fill" style="width:{pct}%;'
                f'background:{_bar_color(pct)}"></div></div><span>{esc(label)}</span></div>')

    rows = []
    for r in servers:
        d = r["data"]
        hd, up = d.get("hardening") or {}, d.get("updates") or {}
        overall = int(d.get("overall_score") or 0)
        hd_label = f"{hd.get('got', 0)}/{hd.get('max', 0)}"
        up_label = f"{up.get('got', 0)}/{up.get('max', 0)}"
        rows.append(f"<tr><td>{esc(r['server_name'])}<br><small>{esc(d.get('hostname', '-'))}</small></td>"
                    f"<td>{esc(d.get('os', '-'))}</td><td>{esc(d.get('generated_at', '-'))}</td>"
                    f"<td>{bar(overall, f'{overall}%')}</td><td>{bar(hd.get('pct'), hd_label)}</td>"
                    f"<td>{bar(up.get('pct'), up_label)}</td></tr>")
    ids: List[Tuple[str, str, str]] = []
    seen = set()
    for r in servers:
        for c in _checks(r["data"]):
            if c.get("id") not in seen:
                seen.add(c.get("id"))
                ids.append((str(c.get("id")), str(c.get("category") or ""), str(c.get("description") or "")))
    badge = {"OK": "badge-ok", "WARN": "badge-warn", "FAIL": "badge-fail"}
    detail = []
    for cid, cat, desc in ids:
        cells = []
        for r in servers:
            m = next((c for c in _checks(r["data"]) if c.get("id") == cid), None)
            cells.append(f'<td><span class="badge {badge.get(str(m.get("status")), "badge-info")}">{esc(m.get("status"))}</span> '
                         f'{esc(m.get("points"))}/{esc(m.get("max"))}</td>' if m else "<td>-</td>")
        detail.append(f"<tr><td>{esc(cat)}</td><td>{esc(desc)}</td>{''.join(cells)}</tr>")
    heads = "".join(f"<th>{esc(r['server_name'])}</th>" for r in servers)
    return f"""<!DOCTYPE html><html lang="es"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>Comparativo de auditorías</title>
<style>
:root{{--ok:#16a34a;--warn:#f59e0b;--fail:#dc2626;--info:#64748b;--bg:#0f172a;--panel:#1e293b;--text:#e2e8f0;--muted:#94a3b8}}
*{{box-sizing:border-box}}body{{margin:0;padding:2rem;background:var(--bg);color:var(--text);font-family:-apple-system,Segoe UI,Roboto,Arial,sans-serif}}
h1{{margin:0 0 .25rem}}p.meta{{color:var(--muted);margin:0 0 1.5rem}}
.panel{{background:var(--panel);border-radius:12px;padding:1.5rem;margin-bottom:1.5rem;border:1px solid #334155}}
table{{width:100%;border-collapse:collapse;font-size:.85rem}}.table-wrap{{overflow-x:auto}}
th,td{{text-align:left;padding:.6rem .7rem;border-bottom:1px solid #334155;white-space:nowrap}}
th{{color:var(--muted);font-weight:600;text-transform:uppercase;font-size:.7rem}}small{{color:var(--muted)}}
.bar-row{{display:flex;align-items:center;gap:.5rem;min-width:160px}}.bar-track{{flex:1;height:12px;background:#334155;border-radius:6px;overflow:hidden}}
.bar-fill{{height:100%;border-radius:6px}}.bar-row span{{width:55px;text-align:right;color:var(--muted);font-size:.8rem}}
.badge{{display:inline-block;padding:.1rem .5rem;border-radius:999px;font-size:.7rem;font-weight:700;color:#0f172a}}
.badge-ok{{background:var(--ok)}}.badge-warn{{background:var(--warn)}}.badge-fail{{background:var(--fail);color:#fff}}.badge-info{{background:var(--info);color:#fff}}
</style></head><body>
<h1>Comparativo de auditorías de servidores</h1><p class="meta">Servidores comparados: <strong>{len(servers)}</strong> · Generado: {now}</p>
<section class="panel"><h2>Resumen por servidor</h2><div class="table-wrap"><table>
<thead><tr><th>Servidor</th><th>SO</th><th>Auditado</th><th>Puntaje general</th><th>Hardening</th><th>Actualizaciones</th></tr></thead>
<tbody>{''.join(rows)}</tbody></table></div></section>
<section class="panel"><h2>Detalle de controles por servidor</h2><div class="table-wrap"><table>
<thead><tr><th>Categoría</th><th>Control</th>{heads}</tr></thead><tbody>{''.join(detail)}</tbody></table></div></section>
</body></html>"""


# ---------------------------------------------------------------- Excel
def _header(ws, columns: int) -> None:
    for cell in ws[1][:columns]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    ws.freeze_panes = "A2"


def comparative_xlsx(results: Sequence[Dict[str, Any]], *, generated_at: datetime) -> bytes:
    servers = _sorted(results)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Servidor", "Hostname", "Sistema operativo", "Auditado", "Puntaje general (%)", "Hardening (%)",
               "Hardening (pts)", "Actualizaciones (%)", "Actualizaciones (pts)", "Sitios web"])
    for r in servers:
        d = r["data"]
        hd, up = d.get("hardening") or {}, d.get("updates") or {}
        ws.append([r["server_name"], d.get("hostname"), d.get("os"), d.get("generated_at"), d.get("overall_score"),
                   hd.get("pct"), f"{hd.get('got', 0)}/{hd.get('max', 0)}", up.get("pct"),
                   f"{up.get('got', 0)}/{up.get('max', 0)}", len(d.get("sites") or [])])
    _header(ws, 10)
    controls = wb.create_sheet("Controles")
    controls.append(["Servidor", "Categoría", "Control", "Estado", "Puntos", "Máximo", "Recomendación"])
    recs = wb.create_sheet("Recomendaciones")
    recs.append(["Servidor", "Estado", "Control", "Recomendación"])
    sites = wb.create_sheet("Sitios")
    sites.append(["Servidor", "Sitio", "Servidor web"])
    for r in servers:
        for c in _checks(r["data"]):
            controls.append([r["server_name"], c.get("category"), c.get("description"), c.get("status"),
                             c.get("points"), c.get("max"), c.get("recommendation")])
        for c in sorted(_checks(r["data"]), key=lambda c: STATUS_ORDER.get(str(c.get("status")), 9)):
            if c.get("status") in ("FAIL", "WARN") and c.get("recommendation"):
                recs.append([r["server_name"], c.get("status"), c.get("description"), c.get("recommendation")])
        for s in r["data"].get("sites") or []:
            if isinstance(s, dict):
                sites.append([r["server_name"], s.get("name"), s.get("source")])
    for sheet, n, widths in ((controls, 7, (22, 14, 50, 9, 8, 8, 90)), (recs, 4, (22, 9, 50, 100)),
                             (sites, 3, (22, 50, 16)), (ws, 10, (22, 24, 30, 20, 12, 12, 12, 14, 14, 10))):
        _header(sheet, n)
        for i, w in enumerate(widths):
            sheet.column_dimensions[chr(65 + i)].width = w
    neutralize_formulas(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- PDF
class _Pagina:
    def __init__(self, generado: str):
        self.generado = generado

    def __call__(self, canvas, doc):
        canvas.saveState()
        ancho, alto = doc.pagesize
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(AZUL_MEDIO)
        canvas.drawRightString(ancho - doc.rightMargin, alto - 12 * mm, TITULO)
        canvas.setStrokeColor(LINEA)
        canvas.line(doc.leftMargin, alto - 13.5 * mm, ancho - doc.rightMargin, alto - 13.5 * mm)
        canvas.setFillColor(GRIS_TEXTO)
        canvas.drawString(doc.leftMargin, 10 * mm, f"Generado: {self.generado}")
        canvas.drawRightString(ancho - doc.rightMargin, 10 * mm, f"Página {doc.page}")
        canvas.restoreState()


def _estado(status: str) -> str:
    color = {"OK": "#1E7B34", "WARN": "#B7791F", "FAIL": "#B42318"}.get(status, "#5B6573")
    return f'<font color="{color}"><b>{_texto(status)}</b></font>'


def comparative_pdf(results: Sequence[Dict[str, Any]], *, generated_at: datetime) -> bytes:
    from reportlab.platypus import Paragraph
    from services.cost_pdf import ESTILOS

    servers = _sorted(results)
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=LETTER, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=22 * mm,
                            bottomMargin=20 * mm, title=TITULO)
    ancho = doc.width
    stamp = generated_at.strftime("%Y-%m-%d %H:%M")
    historia: List[Any] = [_p(TITULO.upper(), "portada_titulo"), Spacer(1, 3 * mm),
                           _p(f"{len(servers)} servidor(es) · generado el {stamp}", "portada_texto"), Spacer(1, 6 * mm)]
    if servers:
        promedio = round(sum(int(r["data"].get("overall_score") or 0) for r in servers) / len(servers))
        fallos = sum(1 for r in servers for c in _checks(r["data"]) if c.get("status") == "FAIL")
        historia.append(_kpis([("Servidores", _texto(len(servers))), ("Puntaje promedio", _texto(f"{promedio}%")),
                               ("Controles en FAIL", _texto(fallos)),
                               ("Mejor", _texto(servers[0]["server_name"]))], ancho))
    historia.append(_p("Resumen por servidor", "h1"))
    filas = [[_p(r["server_name"]), _p(r["data"].get("os") or "—"), _p(f"{r['data'].get('overall_score', 0)}%", "celda_der"),
              _p(f"{(r['data'].get('hardening') or {}).get('pct', 0)}%", "celda_der"),
              _p(f"{(r['data'].get('updates') or {}).get('pct', 0)}%", "celda_der"),
              _p(len(r["data"].get("sites") or []), "celda_der")] for r in servers]
    historia.append(_tabla(["Servidor", "Sistema operativo", "General", "Hardening", "Actualiz.", "Sitios"],
                           filas or [[_p("Sin auditorías"), _p(""), _p(""), _p(""), _p(""), _p("")]],
                           [ancho * 0.22, ancho * 0.34, ancho * 0.11, ancho * 0.11, ancho * 0.11, ancho * 0.11]))
    for r in servers:
        d = r["data"]
        historia.append(PageBreak())
        historia.append(_p(f"{r['server_name']} — {d.get('hostname') or ''}", "h1"))
        historia.append(_p(f"{d.get('os') or 'SO desconocido'} · auditado {d.get('generated_at') or '—'}", "intro"))
        historia.append(_kpis([("General", _texto(f"{d.get('overall_score', 0)}%")),
                               ("Hardening", _texto(f"{(d.get('hardening') or {}).get('got', 0)}/{(d.get('hardening') or {}).get('max', 0)}")),
                               ("Actualizaciones", _texto(f"{(d.get('updates') or {}).get('got', 0)}/{(d.get('updates') or {}).get('max', 0)}"))], ancho))
        controles = sorted(_checks(d), key=lambda c: STATUS_ORDER.get(str(c.get("status")), 9))
        filas = [[Paragraph(_estado(str(c.get("status"))), ESTILOS["celda"]), _p(c.get("category")), _p(c.get("description")),
                  _p(f"{c.get('points')}/{c.get('max')}", "celda_der")] for c in controles]
        historia.append(_p("Controles", "h2"))
        historia.append(_tabla(["Estado", "Categoría", "Control", "Puntos"], filas,
                               [ancho * 0.11, ancho * 0.16, ancho * 0.6, ancho * 0.13]))
        recs = [f"[{c.get('status')}] {c.get('description')}: {c.get('recommendation')}" for c in controles
                if c.get("status") in ("FAIL", "WARN") and c.get("recommendation")]
        if recs:
            historia.append(KeepTogether([_p("Recomendaciones", "h2"), *_vinetas(recs[:1])]))
            historia += _vinetas(recs[1:])
        sitios = [s for s in d.get("sites") or [] if isinstance(s, dict)]
        if sitios:
            historia.append(_p(f"Sitios web ({len(sitios)})", "h2"))
            historia.append(_tabla(["Sitio", "Servidor web"], [[_p(s.get("name")), _p(s.get("source"))] for s in sitios[:200]],
                                   [ancho * 0.75, ancho * 0.25]))
    historia.append(Spacer(1, 4 * mm))
    historia.append(_p("Datos generados por server-audit.sh (solo lectura) en cada servidor. Contiene información "
                       "sensible de infraestructura: trátalo como confidencial.", "nota"))
    pagina = _Pagina(stamp)
    doc.build(historia, onFirstPage=pagina, onLaterPages=pagina)
    return out.getvalue()
