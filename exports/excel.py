# coding: utf-8
"""Generación de Excel del inventario y protección contra Formula Injection.

openpyxl guarda como FÓRMULA cualquier texto que empiece por "=". Además, un
texto que empiece por "=", "+", "-", "@", tabulador o retorno de carro puede
ejecutarse si el usuario edita la celda o la exporta a CSV y la reabre.

Defensa aplicada (sin alterar el dato visible):
  1. la celda se fuerza a tipo texto (``data_type = "s"``): nunca se escribe <f>;
  2. se activa ``quotePrefix`` en su estilo: Excel la trata como texto literal
     también al editarla (equivale a escribir un apóstrofo delante).
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, Iterable, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
MAX_SHEET_TITLE = 31
DETAIL_SCAN_ROWS = 500
WIDTH_SCAN_ROWS = 200
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Inter")
ZEBRA_FILL = PatternFill("solid", fgColor="F9FAFB")
_SIDE = Side(style="thin", color="E5E7EB")
BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
WARNING_HEADERS = ["Servicio", "HTTP Status", "Request ID", "Error Code", "Mensaje"]


def is_formula_like(value: Any) -> bool:
    return isinstance(value, str) and value.startswith(FORMULA_PREFIXES)


def neutralize_formulas(workbook: Workbook) -> int:
    """Convierte en texto literal toda celda con aspecto de fórmula. Devuelve cuántas."""
    count = 0
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if is_formula_like(cell.value):
                    cell.data_type = "s"
                    cell.quotePrefix = True
                    count += 1
    return count


def safe_sheet_title(title: str, used: Iterable[str]) -> str:
    """Título válido para Excel (sin []:*?/\\, ≤31 caracteres, único)."""
    base = _INVALID_SHEET_CHARS.sub("_", title or "").strip("' ") or "Hoja"
    base = base[:MAX_SHEET_TITLE]
    used_lower = {u.lower() for u in used}
    candidate, n = base, 2
    while candidate.lower() in used_lower:
        suffix = f" ({n})"
        candidate = base[: MAX_SHEET_TITLE - len(suffix)] + suffix
        n += 1
    return candidate


def _style_header(worksheet: Worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = BORDER
    worksheet.row_dimensions[1].height = 22
    worksheet.freeze_panes = "A2"


def _detail_columns(columns: Sequence[str], rows: Sequence[Dict[str, Any]]) -> List[str]:
    extra: List[str] = []
    for row in rows[:DETAIL_SCAN_ROWS]:
        detail = row.get("_detalle")
        if isinstance(detail, dict):
            for key in detail:
                if key not in extra and key not in columns:
                    extra.append(key)
    return extra


def _cell_value(row: Dict[str, Any], column: str) -> Any:
    value = row.get(column)
    detail = row.get("_detalle") if isinstance(row.get("_detalle"), dict) else {}
    if (value is None or value == "") and column in detail:
        value = detail[column]
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return str(value)
    return value


def _write_table(workbook: Workbook, title: str, columns: Sequence[str],
                 rows: Sequence[Dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title=safe_sheet_title(title, workbook.sheetnames))
    columns = [c for c in columns if c != "_detalle"]
    columns += _detail_columns(columns, rows)
    worksheet.append(columns)
    _style_header(worksheet)

    for index, row in enumerate(rows, start=2):
        worksheet.append([_cell_value(row, c) for c in columns])
        for cell in worksheet[index]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if index % 2 == 0:
                cell.fill = ZEBRA_FILL

    for col_index, column in enumerate(columns, start=1):
        length = max([len(str(column))] + [len(str(r.get(column, ""))) for r in rows[:WIDTH_SCAN_ROWS]])
        worksheet.column_dimensions[get_column_letter(col_index)].width = min(max(length + 2, 12), 50)
    worksheet.auto_filter.ref = worksheet.dimensions


def _write_warnings(workbook: Workbook, warnings: Sequence[Dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title=safe_sheet_title("Avisos", workbook.sheetnames))
    worksheet.append(WARNING_HEADERS)
    _style_header(worksheet)
    for warning in warnings:
        worksheet.append([
            warning.get("servicio", ""), warning.get("http_status", ""), warning.get("request_id", ""),
            warning.get("error_code", ""), warning.get("mensaje", ""),
        ])
        for cell in worksheet[worksheet.max_row]:
            cell.border = BORDER
    for col_index in range(1, len(WARNING_HEADERS) + 1):
        worksheet.column_dimensions[get_column_letter(col_index)].width = 30


def build_inventory_workbook(tables: Sequence[Dict[str, Any]],
                             warnings: Sequence[Dict[str, Any]] = ()) -> Workbook:
    """Una hoja por tabla (columnas visibles + detalle) y hoja "Avisos" opcional."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, table in enumerate(tables):
        _write_table(workbook, table.get("titulo") or f"Tabla {index + 1}",
                     list(table.get("columnas") or []), list(table.get("filas") or []))
    if warnings:
        _write_warnings(workbook, warnings)
    neutralize_formulas(workbook)
    return workbook


def workbook_bytes(workbook: Workbook) -> io.BytesIO:
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def safe_filename_part(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", value or "")[:60] or "na"
