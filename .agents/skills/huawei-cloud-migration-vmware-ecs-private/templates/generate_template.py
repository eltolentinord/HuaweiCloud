#!/usr/bin/env python3
"""
generate_template.py — 生成 Excel 迁移参数模板 (2 Sheet 页)

运行: python3 generate_template.py
输出: migration_params_template.xlsx

Sheet1: 主机信息 (源端主机, 13 列)
Sheet2: 代理主机信息 (跳板机, 6 列)

注意: AK/SK 不再通过 Excel 提供, 请使用 credential_manager RSA 加密流程。
"""

import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl -q")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter


def _setup_styles():
    """统一样式定义"""
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2196F3", end_color="2196F3", fill_type="solid")
    required_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    optional_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    return header_font, header_fill, required_fill, optional_fill, center_align, thin_border


def _write_sheet(ws, columns, examples, header_font, header_fill, required_fill,
                 optional_fill, center_align, thin_border):
    """写入一个 Sheet 页"""
    # 表头
    for col_idx, name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # 示例行
    for col_idx, value in enumerate(examples, 1):
        cell = ws.cell(row=2, column=col_idx, value=value)
        cell.alignment = center_align
        cell.border = thin_border

    # 列宽
    for col_idx, name in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(name) + 4, 18)


def generate_template(output_path: str = "migration_params_template.xlsx"):
    """生成 Excel 模板 (2 Sheet 页)"""
    wb = Workbook()

    header_font, header_fill, required_fill, optional_fill, center_align, thin_border = _setup_styles()

    # ── Sheet1: 主机信息 (13 列) ──
    ws1 = wb.active
    ws1.title = "主机信息"

    host_columns = [
        "主机名", "内网IP", "端口号", "用户名", "密码",
        "region_id", "region_name", "project_id", "project_name",
        "os_type", "use_public_ip", "target_image_id", "target_AZ",
    ]
    host_examples = [
        "web-server-01", "192.168.1.100", 22, "root", "***",
        "cn-north-4", "华北-北京四", "0axxxxxx", "xxx-project",
        "Linux", "FALSE", "", "",
    ]
    _write_sheet(ws1, host_columns, host_examples, header_font, header_fill,
                 required_fill, optional_fill, center_align, thin_border)

    # ── Sheet2: 代理主机信息 (6 列) ──
    ws2 = wb.create_sheet("代理主机信息")

    proxy_columns = ["主机名称", "公网IP", "私网IP", "端口", "用户名", "密码"]
    proxy_examples = ["proxy-ecs-01", "1.2.3.4", "10.0.1.10", 22, "root", "***"]
    _write_sheet(ws2, proxy_columns, proxy_examples, header_font, header_fill,
                 required_fill, optional_fill, center_align, thin_border)

    # 保存
    wb.save(output_path)
    print(f"✓ Template generated: {output_path}")
    print(f"  Sheet1: 主机信息 ({len(host_columns)} 列)")
    print(f"  Sheet2: 代理主机信息 ({len(proxy_columns)} 列)")
    print("  注意: AK/SK 请通过 credential_manager RSA 加密流程提供, 不再使用 Excel Sheet")


if __name__ == "__main__":
    output = sys.argv[1] if len(sys.argv) > 1 else "migration_params_template.xlsx"
    generate_template(output)
