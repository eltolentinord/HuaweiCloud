# coding: utf-8
"""Calculadora de precios: ítems ECS/EVS/EIP, modos mensual/anual/por uso, lista con total y
exportación. Todo simulado: sin llamadas a Huawei Cloud."""

import csv
import io
import unittest
from datetime import datetime, timezone
from decimal import Decimal
from types import SimpleNamespace
from unittest import mock

from openpyxl import load_workbook

from tests.helpers import FAKE_AK, FAKE_SK, api_error
from tests.test_cost_compare import HK, MX, SG, CostCompareApiTestCase

from core.authz import Principal
from costs import huawei_pricing
from costs.calculator import item_from_payload, price_item, price_list, quotable
from costs.catalog import Flavor, PriceCatalog, Quote
from costs.configuration import Component, ConfigurationError
from db.session import session_scope
from routers.security import get_principal

NOW = datetime(2026, 10, 5, 12, 0, tzinfo=timezone.utc)
REGIONS = [SG, MX, HK]


def q(region, product, spec, size, amount, mode="monthly", period=None):
    period = period or {"monthly": "month", "yearly": "year", "on_demand": "hour"}[mode]
    return Quote(region=region, product=product, spec=spec, size=size, billing_mode=mode, amount=Decimal(amount),
                 period=period, currency="USD", fetched_at=NOW, source_detail="test")


def ecs_item(mode="monthly", duration=1, quantity=1, eip=None, data_disks=(), region=MX, flavor="x1.2u.4g"):
    return {"region": region, "billing_mode": mode, "duration": duration, "quantity": quantity, "product": "ecs",
            "config": {"flavor": flavor, "os_type": "linux", "system_disk": {"volume_type": "GPSSD", "size_gb": 40},
                       "data_disks": list(data_disks), "eip": eip}}


class TestItems(unittest.TestCase):
    def test_ecs_components_and_bss_codes(self):
        item = item_from_payload(ecs_item(eip={"ip_type": "5_bgp", "bandwidth_mode": "bandwidth", "bandwidth_mbps": 10},
                                          data_disks=[{"volume_type": "SSD", "size_gb": 100}]), REGIONS)
        self.assertEqual([(c.product, c.spec, c.size) for c in item.components()], [
            ("ecs", "x1.2u.4g.linux", 0), ("evs", "GPSSD", 40), ("evs", "SSD", 100), ("ip", "5_bgp", 0),
            ("bandwidth", "19_bgp", 10)])
        traffic = item_from_payload({"region": MX, "billing_mode": "monthly", "duration": 2, "quantity": 1,
                                     "product": "eip", "config": {"ip_type": "5_sbgp", "bandwidth_mode": "traffic",
                                                                  "bandwidth_mbps": 50, "traffic_gb": 300}}, REGIONS)
        self.assertEqual([(c.product, c.spec, c.size) for c in traffic.components()],
                         [("ip", "5_sbgp", 0), ("traffic", "12_sbgp", 0)])
        self.assertEqual(traffic.traffic_gb, 300)
        self.assertEqual(quotable([traffic])[(MX, "on_demand")][0].product, "traffic")   # tráfico: siempre por uso
        self.assertEqual(quotable([traffic])[(MX, "monthly")][0].product, "ip")

    def test_validation(self):
        bad = [dict(ecs_item(), region="xx-1"), dict(ecs_item(), billing_mode="weekly"), ecs_item(duration=10),
               ecs_item(mode="yearly", duration=4), ecs_item(quantity=0), dict(ecs_item(), product="rds"),
               ecs_item(flavor="BAD!"), {"region": MX, "product": "evs", "config": {"volume_type": "X", "size_gb": 10}}]
        for payload in bad:
            with self.subTest(payload=payload), self.assertRaises(ConfigurationError):
                item_from_payload(payload, REGIONS)


class TestPricing(CostCompareApiTestCase):
    def setUp(self):
        super().setUp()
        with session_scope(self.factory) as session:
            store = PriceCatalog(session)
            store.store([q(MX, "ecs", "x1.2u.4g.linux", 0, "30.00"), q(MX, "evs", "GPSSD", 40, "4.00"),
                         q(MX, "ecs", "x1.2u.4g.linux", 0, "300.00", "yearly"), q(MX, "evs", "GPSSD", 40, "40.00", "yearly"),
                         q(MX, "ecs", "x1.2u.4g.linux", 0, "0.05", "on_demand"), q(MX, "evs", "GPSSD", 40, "0.01", "on_demand"),
                         q(MX, "ip", "5_bgp", 0, "3.00"), q(MX, "traffic", "12_bgp", 0, "0.10", "on_demand", "gb")])
            store.replace_flavors(MX, [Flavor("x1.2u.4g", 2, 4096, "computingv3", "x1", status="normal")], fetched_at=NOW)
        self.calc = f"/api/clients/{self.cid}/accounts/{self.aid}/calculator"

    def priced(self, payload):
        with session_scope(self.factory) as session:
            return price_item(PriceCatalog(session), item_from_payload(payload, REGIONS), has_project=True)

    def test_monthly_yearly_and_on_demand(self):
        self.assertEqual(self.priced(ecs_item(duration=3, quantity=2))["total"], "204.00")      # (30+4)×3×2
        yearly = self.priced(ecs_item(mode="yearly", duration=2))
        self.assertEqual(yearly["total"], "680.00")                                              # anual oficial ×2
        self.assertEqual(yearly["lines"][0]["unit"], "año")
        self.assertEqual(self.priced(ecs_item(mode="on_demand", duration=730))["total"], "43.80")  # (0.05+0.01)×730

    def test_traffic_uses_estimated_gb_and_missing_price_blocks_total(self):
        eip = {"ip_type": "5_bgp", "bandwidth_mode": "traffic", "bandwidth_mbps": 5, "traffic_gb": 200}
        priced = self.priced({"region": MX, "billing_mode": "monthly", "duration": 1, "quantity": 1, "product": "eip",
                              "config": eip})
        self.assertEqual(priced["total"], "23.00")                                               # 3 + 0.10×200
        missing = self.priced(ecs_item(mode="yearly", eip={"ip_type": "5_bgp", "bandwidth_mode": "bandwidth",
                                                             "bandwidth_mbps": 5}))
        self.assertIsNone(missing["total"])
        self.assertIn("Precio no disponible", missing["reason"])

    def test_list_totals_only_complete_items(self):
        with session_scope(self.factory) as session:
            catalog = PriceCatalog(session)
            items = [price_item(catalog, item_from_payload(p, REGIONS), has_project=True)
                     for p in (ecs_item(), ecs_item(region=SG))]
        result = price_list(items)
        self.assertEqual((result["totals"], result["complete"], result["incomplete"]),
                         ([{"currency": "USD", "total": "34.00"}], 1, 1))

    # -------------------------------------------------------------- API
    def test_options_with_flavor_prices(self):
        body = self.call("GET", f"{self.calc}/options", params={"region": MX}).json()
        self.assertTrue(body["has_project"])
        self.assertEqual(body["flavors"][0]["price"], {"amount": "30.00000000", "currency": "USD", "period": "month"})
        self.assertFalse(next(r for r in body["regions"] if r["id"] == HK)["has_project"])
        self.assertEqual([m["id"] for m in body["modes"]], ["monthly", "yearly", "on_demand"])

    def test_quote_list_and_server_side_recalculation(self):
        quote = self.call("POST", f"{self.calc}/quote", json=ecs_item()).json()
        self.assertEqual(quote["total"], "34.00")
        tampered = dict(ecs_item(), total="0.01", lines=[])
        listed = self.call("POST", f"{self.calc}/list", json={"items": [tampered, ecs_item(region=HK)]}).json()
        self.assertEqual(listed["totals"], [{"currency": "USD", "total": "34.00"}])
        hk = listed["items"][1]
        self.assertFalse(hk["has_project"])
        self.assertTrue(any("Sin Project" in w for w in hk["warnings"]))
        self.assertEqual(self.call("POST", f"{self.calc}/quote", json=ecs_item(duration=99)).status_code, 422)

    def test_exports(self):
        body = {"items": [ecs_item(quantity=2), ecs_item(region=HK)]}
        xlsx = self.call("POST", f"{self.calc}/export", params={"format": "xlsx"}, json=body)
        self.assertEqual(xlsx.status_code, 200, xlsx.text)
        wb = load_workbook(io.BytesIO(xlsx.content))
        self.assertEqual(wb.sheetnames, ["Resumen", "Detalle"])
        values = [c.value for row in wb["Resumen"].iter_rows() for c in row]
        self.assertIn("68.00", values)
        self.assertIn("Precio no disponible", values)
        pdf = self.call("POST", f"{self.calc}/export", params={"format": "pdf"}, json=body)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertEqual(pdf.headers["content-type"], "application/pdf")
        rows = list(csv.reader(io.StringIO(self.call("POST", f"{self.calc}/export", params={"format": "csv"},
                                                     json=body).content.decode("utf-8-sig"))))
        self.assertEqual(rows[0][:3], ["#", "Producto", "Región"])
        self.assertEqual(rows[-1][7], "68.00")

    def test_refresh_prices_yearly_and_traffic(self):
        seen = []

        class FakeBss:
            def list_rate_on_period_detail(self, request):
                info = request.body.product_infos[0]
                seen.append(("period", info.period_type, info.resource_spec))
                return SimpleNamespace(currency="USD", official_website_rating_result=SimpleNamespace(
                    official_website_amount=Decimal("100.00"), measure_id=1))

            def list_on_demand_resource_ratings(self, request):
                info = request.body.product_infos[0]
                seen.append(("demand", info.usage_factor, info.resource_spec))
                return SimpleNamespace(currency="USD", official_website_amount=Decimal("0.09"), measure_id=1)

        item = {"region": MX, "billing_mode": "yearly", "duration": 1, "quantity": 1, "product": "eip",
                "config": {"ip_type": "5_bgp", "bandwidth_mode": "traffic", "bandwidth_mbps": 5, "traffic_gb": 10}}
        with mock.patch("costs.huawei_pricing.default_bss_builder", lambda clients: FakeBss()):
            response = self.call("POST", f"{self.calc}/prices/refresh", json={"items": [item, ecs_item(region=HK)]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(("period", 3, "5_bgp"), seen)                 # anual: period_type 3
        self.assertIn(("demand", "upflow", "12_bgp"), seen)         # tráfico: por uso, upflow
        self.assertTrue(all(f["region"] == HK and "Sin Project" in f["reason"] for f in response.json()["failures"]))
        priced = self.call("POST", f"{self.calc}/quote", json=item).json()
        self.assertEqual(priced["total"], "100.90")                 # 100 (IP 1 año) + 0.09 × 10 GB

    def test_refresh_errors_are_safe_and_role_checked(self):
        class Denied:
            def list_rate_on_period_detail(self, request):
                raise api_error(403, f"Access denied {FAKE_AK} {FAKE_SK}", "CBC.0151")

        with mock.patch("costs.huawei_pricing.default_bss_builder", lambda clients: Denied()):
            response = self.call("POST", f"{self.calc}/prices/refresh", json={"items": [ecs_item()]})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(FAKE_AK, response.text)
        self.assertNotIn(FAKE_SK, response.text)
        self.app.dependency_overrides[get_principal] = lambda: Principal(
            subject="v", kind="user", client_roles={str(self.cid): "viewer"})
        self.assertEqual(self.call("POST", f"{self.calc}/prices/refresh", json={"items": [ecs_item()]}).status_code, 403)
        self.assertEqual(self.call("POST", f"{self.calc}/quote", json=ecs_item()).status_code, 200)

    def test_isolation(self):
        other = f"/api/clients/{self.other_cid}/accounts/{self.aid}/calculator/quote"
        self.assertEqual(self.call("POST", other, json=ecs_item()).status_code, 404)


def obs_item(storage_class="standard", mode="on_demand", duration=730):
    return {"region": MX, "billing_mode": mode, "duration": duration, "quantity": 1, "product": "obs",
            "config": {"storage_class": storage_class}}


class TestObsAndRds(CostCompareApiTestCase):
    """OBS usa los códigos oficiales de Huawei; RDS no se puede cotizar por API."""

    def setUp(self):
        super().setUp()
        self.calc = f"/api/clients/{self.cid}/accounts/{self.aid}/calculator"

    def test_obs_uses_official_codes_and_is_on_demand_only(self):
        item = item_from_payload(obs_item("warm"), REGIONS)
        self.assertEqual([(c.product, c.spec, c.size) for c in item.components()],
                         [("obs_storage", "obs.warm", 0)])
        self.assertEqual(huawei_pricing.PRODUCT_CODES["obs_storage"],
                         ("hws.service.type.obs", "hws.resource.type.obs", None))
        for bad in (obs_item(mode="monthly", duration=1), obs_item(mode="yearly", duration=1),
                    obs_item(storage_class="GLACIER")):
            with self.subTest(bad=bad["config"]), self.assertRaises(ConfigurationError):
                item_from_payload(bad, REGIONS)
        self.assertEqual(self.call("POST", f"{self.calc}/quote", json=obs_item(mode="monthly", duration=1)).status_code, 422)

    def test_obs_quote_is_hourly_without_size(self):
        seen = []

        class FakeBss:
            def list_on_demand_resource_ratings(self, request):
                info = request.body.product_infos[0]
                seen.append((info.cloud_service_type, info.resource_type, info.resource_spec, info.usage_factor,
                             info.usage_measure_id, getattr(info, "resource_size", None)))
                return SimpleNamespace(currency="USD", official_website_amount=Decimal("0.02"), measure_id=1)

        with mock.patch("costs.huawei_pricing.default_bss_builder", lambda clients: FakeBss()):
            response = self.call("POST", f"{self.calc}/prices/refresh", json={"items": [obs_item()]})
        self.assertEqual(response.json()["stored"], 1)
        self.assertEqual(seen, [("hws.service.type.obs", "hws.resource.type.obs", "obs.standard",
                                 "Duration", 4, None)])      # plantilla oficial: Duration, 1 h, sin tamaño
        priced = self.call("POST", f"{self.calc}/quote", json=obs_item()).json()
        self.assertEqual(priced["total"], "14.60")            # 0.02 × 730 h
        self.assertTrue(any("por hora" in w for w in priced["warnings"]))

    def test_rds_is_not_a_calculator_product(self):
        rds = {"region": MX, "billing_mode": "monthly", "duration": 1, "quantity": 1, "product": "rds", "config": {}}
        with self.assertRaises(ConfigurationError):
            item_from_payload(rds, REGIONS)
        self.assertEqual(self.call("POST", f"{self.calc}/quote", json=rds).status_code, 422)
        self.assertNotIn("rds", self.call("GET", f"{self.calc}/options").json())

    def test_rds_catalog_stays_informational_without_prices(self):
        class FakeRds:
            def list_datastores(self, request):
                return SimpleNamespace(data_stores=[SimpleNamespace(name="8.0")])

            def list_flavors(self, request):
                return SimpleNamespace(flavors=[SimpleNamespace(
                    vcpus="2", ram=4, spec_code="rds.mysql.n1.large.2.ha", instance_mode="ha",
                    az_status={"mx-a": "normal"})])

        with mock.patch("costs.huawei_pricing.default_rds_builder", lambda c, r, p: FakeRds()):
            self.assertEqual(self.call("POST", f"{self.calc}/rds/refresh",
                                       params={"region": MX, "engine": "MySQL"}).status_code, 200)
        body = self.call("GET", f"{self.calc}/rds/flavors", params={"region": MX, "engine": "MySQL"}).json()
        self.assertFalse(body["pricing_supported"])
        self.assertNotIn("price", body["flavors"][0])
        self.assertEqual(body["flavors"][0]["spec_code"], "rds.mysql.n1.large.2.ha")


class TestQuoteComponentModes(unittest.TestCase):
    def test_yearly_period_type_and_unknown_mode(self):
        captured = {}

        class Fake:
            def list_rate_on_period_detail(self, request):
                captured["type"] = request.body.product_infos[0].period_type
                return SimpleNamespace(currency="USD", official_website_rating_result=SimpleNamespace(
                    official_website_amount=Decimal("12"), measure_id=1))

        quote = huawei_pricing.quote_component(Fake(), project_id="p", region=MX, billing_mode="yearly",
                                               component=Component("ip", "5_bgp"))
        self.assertEqual((captured["type"], quote.period, quote.billing_mode), (3, "year", "yearly"))
        with self.assertRaises(huawei_pricing.PricingDataError):
            huawei_pricing.quote_component(Fake(), project_id="p", region=MX, billing_mode="weekly",
                                           component=Component("ip", "5_bgp"))


class TestCalculatorPage(unittest.TestCase):
    def test_page_and_sidebar(self):
        from fastapi.testclient import TestClient
        from app import app
        http = TestClient(app)
        page = http.get("/costos/calculadora").text
        for marker in ('data-product="ecs"', 'id="ecsFlavors"', 'id="list"', 'data-export="pdf"',
                       'src="/static/calculator.js?v='):
            self.assertIn(marker, page)
        self.assertIn('href="/costos/calculadora"', http.get("/clientes").text)
        script = http.get("/static/calculator.js").text
        for marker in ("/calculator", "/prices/refresh", "function esc("):
            self.assertIn(marker, script)


if __name__ == "__main__":
    unittest.main()
