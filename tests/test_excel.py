# coding: utf-8
"""Exportación Excel: formato conservado y protección contra Formula Injection."""

import io
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from tests.helpers import ROOT  # noqa: F401  (asegura sys.path)

from exports.excel import (
    build_inventory_workbook,
    is_formula_like,
    neutralize_formulas,
    safe_sheet_title,
    workbook_bytes,
)
from mock_data_MOCK import MOCK_RESPONSE
from services.cost_excel import generar_excel_costos

PAYLOADS = ["=1+1", "+SUM(A1:A2)", "-2+3", "@SUM(A1)", "=HYPERLINK(\"http://x\",\"y\")",
            "\t=cmd", "=cmd|' /C calc'!A0"]


def sheet_xml(buffer: io.BytesIO) -> str:
    with zipfile.ZipFile(buffer) as archive:
        return "".join(archive.read(n).decode("utf-8") for n in archive.namelist()
                       if n.startswith("xl/worksheets/"))


class TestFormulaInjection(unittest.TestCase):
    def build(self):
        rows = [{"Nombre": p, "Estado": "ACTIVE", "_detalle": {"Notas": p}} for p in PAYLOADS]
        tables = [{"titulo": "ECS", "columnas": ["Nombre", "Estado"], "filas": rows}]
        warnings = [{"servicio": "=HSS", "http_status": "403", "mensaje": "+evil"}]
        return build_inventory_workbook(tables, warnings)

    def test_no_formula_elements_written(self):
        xml = sheet_xml(workbook_bytes(self.build()))
        self.assertNotIn("<f>", xml)
        self.assertNotIn("<f ", xml)

    def test_values_preserved_as_text(self):
        workbook = load_workbook(workbook_bytes(self.build()))
        ws = workbook["ECS"]
        names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(names, PAYLOADS)
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=1)
            self.assertEqual(cell.data_type, "s")
            self.assertTrue(cell.quotePrefix)
        self.assertEqual(workbook["Avisos"]["A2"].value, "=HSS")
        self.assertTrue(workbook["Avisos"]["A2"].quotePrefix)

    def test_detection(self):
        for value in PAYLOADS:
            self.assertTrue(is_formula_like(value), value)
        for value in ("web-01", "", None, 5, -3.2, "192.168.0.1", "ñandú"):
            self.assertFalse(is_formula_like(value), value)

    def test_numbers_are_untouched(self):
        tables = [{"titulo": "T", "columnas": ["n"], "filas": [{"n": -5}, {"n": 3.5}]}]
        ws = load_workbook(workbook_bytes(build_inventory_workbook(tables)))["T"]
        self.assertEqual((ws["A2"].value, ws["A3"].value), (-5, 3.5))
        self.assertEqual(ws["A2"].data_type, "n")

    def test_cost_excel_is_protected_too(self):
        analysis = {"monedas": [{
            "currency": "USD", "total_mes_a": 1.0, "total_mes_b": 2.0, "diferencia": 1.0, "variacion_txt": "100.00%",
            "region_mayor_gasto": "=evil", "producto_mas_costoso": "+evil", "producto_mas_aumento": "-",
            "producto_mas_reduccion": "@x", "por_producto": [], "por_region": [], "por_region_producto": [],
            "mayores_aumentos": [], "mayores_reducciones": [], "servicios_nuevos": [], "servicios_eliminados": [],
        }]}
        with TemporaryDirectory() as tmp:
            path = generar_excel_costos(analysis, "2026-01", "2026-02", Path(tmp))
            with open(path, "rb") as fh:
                self.assertNotIn("<f>", sheet_xml(io.BytesIO(fh.read())))


class TestFormatPreserved(unittest.TestCase):
    def test_mock_tables_export_with_format(self):
        tables = MOCK_RESPONSE["tables"]
        workbook = build_inventory_workbook(tables, [{"servicio": "HSS", "mensaje": "aviso"}])
        self.assertEqual(workbook.sheetnames[-1], "Avisos")
        self.assertEqual(len(workbook.sheetnames), len(tables) + 1)
        ws = workbook["ECS"]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.auto_filter.ref)
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[:12], tables[0]["columnas"])
        self.assertIn("ECS ID", headers)  # columnas de detalle añadidas
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["A1"].fill.fgColor.rgb[-6:], "1D4ED8")
        self.assertEqual(workbook["Avisos"].freeze_panes, "A2")

    def test_no_warnings_sheet_when_empty(self):
        workbook = build_inventory_workbook([{"titulo": "X", "columnas": ["a"], "filas": [{"a": 1}]}])
        self.assertEqual(workbook.sheetnames, ["X"])

    def test_sheet_titles_are_sanitized_and_unique(self):
        self.assertEqual(safe_sheet_title("a/b:c*d?[e]", []), "a_b_c_d__e_")
        self.assertEqual(safe_sheet_title("ECS", ["ECS"]), "ECS (2)")
        self.assertLessEqual(len(safe_sheet_title("x" * 50, ["x" * 31])), 31)

    def test_neutralize_counts(self):
        workbook = build_inventory_workbook([{"titulo": "T", "columnas": ["a"], "filas": [{"a": "ok"}]}])
        workbook["T"]["B2"] = "=1+1"
        self.assertEqual(neutralize_formulas(workbook), 1)


if __name__ == "__main__":
    unittest.main()
