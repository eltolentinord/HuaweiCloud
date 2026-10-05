# coding: utf-8
"""Comparación de costos por recurso, hallazgos y exportación PDF/Excel de ``/costos``.
Todo simulado: sin llamadas a Huawei Cloud."""

import io
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

import huaweicloudsdkbss.v2 as bss_china
import huaweicloudsdkbssintl.v2 as bss_intl
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.helpers import FAKE_AK, FAKE_SK, api_error
from tests.test_bss_site import INTL_ENDPOINT, FakeBuilder

from app import app
from services import huawei_costs
from services.cost_analysis import analizar, analizar_recursos, hallazgos, nombre_mes, nombre_servicio
from services.cost_excel import generar_excel_costos
from services.cost_pdf import generar_pdf_costos

S = "hws.service.type."
COSTOS_A = [{"region": "na-mexico-1", "product": S + "ec2", "currency": "USD", "amount": 100.0},
            {"region": "la-north-2", "product": S + "vpn", "currency": "USD", "amount": 20.0},
            {"region": "ap-southeast-1", "product": S + "modelarts", "currency": "USD", "amount": 8.89}]
COSTOS_B = [{"region": "na-mexico-1", "product": S + "ec2", "currency": "USD", "amount": 90.0},
            {"region": "la-north-2", "product": S + "vpn", "currency": "USD", "amount": 240.0},
            {"region": "la-north-2", "product": S + "ebs", "currency": "USD", "amount": 60.0}]


def recurso(rid, nombre, servicio, region, importe, moneda="USD"):
    return {"resource_id": rid, "resource_name": nombre, "region": region, "region_name": "",
            "service_code": S + servicio, "service_name": "", "resource_type_name": "", "spec": "spec-1",
            "amount": importe, "currency": moneda}


RECURSOS_A = [recurso("vm-1", "web-01", "ec2", "na-mexico-1", 60.0), recurso("vm-1", "web-01", "ec2", "na-mexico-1", 40.0),
              recurso("vpn-1", "vpngw", "vpn", "la-north-2", 20.0), recurso("old-1", "retirado", "ec2", "na-mexico-1", 5.0)]
RECURSOS_B = [recurso("vm-1", "web-01", "ec2", "na-mexico-1", 90.0), recurso("vpn-1", "vpngw", "vpn", "la-north-2", 240.0),
              recurso("vol-1", "<script>alert(1)</script>", "ebs", "la-north-2", 60.0)]


def analisis_completo(recursos=True):
    analisis = analizar(COSTOS_A, COSTOS_B, "2026-08", "2026-09")
    if recursos:
        analisis["recursos"] = analizar_recursos(RECURSOS_A, RECURSOS_B)
    analisis["hallazgos"] = hallazgos(analisis, "2026-08", "2026-09")
    return analisis


class TestNombres(unittest.TestCase):
    def test_known_unknown_and_empty_service_codes(self):
        self.assertEqual(nombre_servicio(S + "ec2"), {"codigo": S + "ec2", "sigla": "ECS", "nombre": "Elastic Cloud Server"})
        self.assertEqual(nombre_servicio(S + "modelarts")["sigla"], "ModelArts")
        self.assertEqual(nombre_servicio(S + "nuevo", "Servicio Nuevo"), {"codigo": S + "nuevo", "sigla": "NUEVO",
                                                                           "nombre": "Servicio Nuevo"})
        self.assertEqual(nombre_servicio(None)["sigla"], "Sin servicio")

    def test_region_names_prefer_catalog_and_skip_non_latin_bss_names(self):
        from services.cost_analysis import nombre_region
        self.assertEqual(nombre_region("ap-southeast-1", {"ap-southeast-1": "中国-香港"}), "CN-Hong Kong")
        self.assertEqual(nombre_region("xx-new-1", {"xx-new-1": "中国-香港"}), "xx-new-1")
        self.assertEqual(nombre_region("xx-new-1", {"xx-new-1": "EU-Paris"}), "EU-Paris")

    def test_pdf_draws_chinese_text_with_an_embedded_cjk_font(self):
        from services import cost_pdf
        self.assertEqual(cost_pdf._texto("x0.4u.8g <b>"), "x0.4u.8g &lt;b&gt;")
        marcado = cost_pdf._texto("deepseek-v3.1 输入Tokens")
        self.assertRegex(marcado, r'deepseek-v3\.1 <font name="[^"]+">输入</font>Tokens')
        analisis = analisis_completo()
        analisis["recursos"]["monedas"][0]["top_mes_b"][0]["spec"] = "deepseek-v3.1 输入Tokens 百万"
        self.assertTrue(generar_pdf_costos(analisis, "2026-08", "2026-09").startswith(b"%PDF"))

    def test_month_names(self):
        self.assertEqual(nombre_mes("2026-07"), "julio 2026")
        self.assertEqual(nombre_mes("raro"), "raro")

    def test_analysis_is_additive_with_readable_names(self):
        m = analizar(COSTOS_A, COSTOS_B, "2026-08", "2026-09")["monedas"][0]
        ecs = next(f for f in m["por_producto"] if f["product"] == S + "ec2")
        self.assertEqual((ecs["servicio_sigla"], ecs["mes_a"], ecs["mes_b"]), ("ECS", 100.0, 90.0))
        mx2 = next(f for f in m["por_region"] if f["region"] == "la-north-2")
        self.assertEqual(mx2["region_nombre"], "LA-Mexico City2")
        detalle = m["por_region_detalle"][0]
        self.assertEqual(detalle["region"], "la-north-2")  # mayor gasto en el mes B primero
        self.assertEqual([s["servicio_sigla"] for s in detalle["servicios"]], ["VPN", "EVS"])


class TestRecursos(unittest.TestCase):
    def test_top_changes_new_and_removed(self):
        r = analizar_recursos(RECURSOS_A, RECURSOS_B)["monedas"][0]
        self.assertEqual((r["recursos_mes_a"], r["recursos_mes_b"]), (3, 3))
        self.assertEqual([f["resource_id"] for f in r["top_mes_b"]], ["vpn-1", "vm-1", "vol-1"])
        self.assertEqual(r["top_mes_a"][0]["mes_a"], 100.0)  # registros del mismo recurso se suman
        self.assertEqual([f["resource_id"] for f in r["mayores_aumentos"]], ["vpn-1", "vol-1"])
        self.assertEqual([f["resource_id"] for f in r["mayores_reducciones"]], ["vm-1", "old-1"])
        self.assertEqual([f["resource_id"] for f in r["nuevos"]], ["vol-1"])
        self.assertEqual([f["resource_id"] for f in r["eliminados"]], ["old-1"])
        self.assertEqual(r["top_mes_b"][0]["region_nombre"], "LA-Mexico City2")
        self.assertNotIn("_a", r["top_mes_b"][0])

    def test_currencies_are_never_mixed(self):
        r = analizar_recursos([recurso("a", "a", "ec2", "x", 1.0, "USD")], [recurso("b", "b", "ec2", "x", 2.0, "EUR")])
        self.assertEqual([m["currency"] for m in r["monedas"]], ["EUR", "USD"])

    def test_empty_inputs(self):
        self.assertEqual(analizar_recursos([], []), {"monedas": []})


class TestHallazgos(unittest.TestCase):
    def test_sentences_come_only_from_data(self):
        h = analisis_completo()["hallazgos"]
        texto = " ".join(h["hallazgos"])
        self.assertIn("El consumo total aumentó de USD 128.89 en agosto 2026 a USD 390.00 en septiembre 2026", texto)
        self.assertIn("LA-Mexico City2 aumentó USD 280.00", texto)
        self.assertIn("VPN fue el servicio con mayor aumento", texto)
        self.assertIn("vpngw (VPN, LA-Mexico City2): USD 240.00", texto)
        self.assertTrue(any("vpngw" in t for t in h["revisar"]))
        self.assertFalse(any("Revisar <script>" in t for t in h["revisar"]))  # nuevo: va en "Confirmar", no repetido

    def test_no_consumption_no_sentences(self):
        vacio = analizar([], [], "2026-08", "2026-09")
        self.assertEqual(hallazgos(vacio, "2026-08", "2026-09", preliminar=True), {"hallazgos": [], "revisar": []})


class TestExports(unittest.TestCase):
    def test_pdf_with_resources_and_long_names(self):
        analisis = analisis_completo()
        analisis["recursos"]["monedas"][0]["top_mes_b"][0]["resource_name"] = "x" * 300
        pdf = generar_pdf_costos(analisis, "2026-08", "2026-09", extraccion="2026-10-05 10:00:00", preliminar=True)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertGreater(len(pdf), 3000)

    def test_pdf_without_data_and_with_resource_error(self):
        self.assertTrue(generar_pdf_costos({"monedas": []}, "2026-08", "2026-09").startswith(b"%PDF"))
        error = {"mensaje": "Access denied.", "http_status": 403, "error_code": "CBC.0150", "request_id": "r-1"}
        pdf = generar_pdf_costos(analisis_completo(recursos=False), "2026-08", "2026-09", recursos_error=error)
        self.assertTrue(pdf.startswith(b"%PDF"))

    def test_excel_adds_resource_sheets_without_removing_existing(self):
        with TemporaryDirectory() as tmp:
            ruta = generar_excel_costos(analisis_completo(), "2026-08", "2026-09", Path(tmp))
            wb = load_workbook(ruta)
        for hoja in ("00_Resumen_Ejecutivo", "02_Costos_por_Producto", "07b_Servicios_Eliminados", "99_Errores",
                     "09_Top_Recursos", "10_Recursos_Aumentos_Reducc", "11_Recursos_Nuevos_Elimin"):
            self.assertIn(hoja, wb.sheetnames)
        self.assertEqual(wb.sheetnames[-1], "99_Errores")
        self.assertEqual(wb["02_Costos_por_Producto"]["C1"].value, "Servicio")
        self.assertIn("ECS · Elastic Cloud Server", [c.value for c in wb["02_Costos_por_Producto"]["C"]])
        nombres = [c.value for c in wb["09_Top_Recursos"]["B"]]
        self.assertIn("vpngw", nombres)

    def test_excel_without_resource_detail_keeps_old_sheets(self):
        with TemporaryDirectory() as tmp:
            wb = load_workbook(generar_excel_costos(analisis_completo(recursos=False), "2026-08", "2026-09", Path(tmp)))
        self.assertNotIn("09_Top_Recursos", wb.sheetnames)


class TestCompareEndpoint(unittest.TestCase):
    def setUp(self):
        self.http = TestClient(app)
        self.body = {"ak": FAKE_AK, "sk": FAKE_SK, "month_a": "2026-08", "month_b": "2026-09"}

    def fake_mes(self, ak, sk, mes, *args):
        return {"records": COSTOS_A if mes == "2026-08" else COSTOS_B, "errors": []}

    def fake_recursos(self, ak, sk, mes):
        return RECURSOS_A if mes == "2026-08" else RECURSOS_B

    def test_compare_includes_resources_and_findings(self):
        with mock.patch("routers.costs.consultar_mes", self.fake_mes), \
                mock.patch("routers.costs.consultar_recursos_mes", self.fake_recursos):
            data = self.http.post("/api/costs/compare", json=self.body).json()
        self.assertFalse(data["error"])
        self.assertIsNone(data["recursos_error"])
        self.assertEqual(data["analysis"]["recursos"]["monedas"][0]["top_mes_b"][0]["resource_name"], "vpngw")
        self.assertTrue(data["analysis"]["hallazgos"]["hallazgos"])

    def test_resource_failure_keeps_comparison_and_is_safe(self):
        def denied(ak, sk, mes):
            raise huawei_costs.CostApiError(f"Access denied {FAKE_AK} {FAKE_SK}", http_status=403,
                                            request_id="req-1", error_code="CBC.0150")
        with mock.patch("routers.costs.consultar_mes", self.fake_mes), \
                mock.patch("routers.costs.consultar_recursos_mes", denied):
            response = self.http.post("/api/costs/compare", json=self.body)
        data = response.json()
        self.assertFalse(data["error"])
        self.assertEqual(data["analysis"]["monedas"][0]["total_mes_b"], 390.0)
        self.assertNotIn("recursos", data["analysis"])
        self.assertEqual((data["recursos_error"]["error_code"], data["recursos_error"]["request_id"]),
                         ("CBC.0150", "req-1"))
        self.assertNotIn(FAKE_AK, response.text)
        self.assertNotIn(FAKE_SK, response.text)

    def test_pdf_endpoint(self):
        analisis = analisis_completo()
        response = self.http.post("/api/costs/export/pdf", json={
            "analysis": analisis, "month_a": "2026-08", "month_b": "2026-09", "extraction_time": "2026-10-05 10:00:00"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertIn('filename="Huawei_Costos_2026-08_vs_2026-09.pdf"', response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertNotIn(FAKE_SK.encode(), response.content)

    def test_pdf_endpoint_rejects_bad_months(self):
        response = self.http.post("/api/costs/export/pdf", json={"analysis": {}, "month_a": "x\r\nevil", "month_b": "2026-09"})
        self.assertEqual(response.status_code, 400)


class TestResourceQueryUsesInternationalBss(unittest.TestCase):
    def test_resource_records_use_bss_international_and_flatten(self):
        class Fake:
            def __init__(self):
                self.requests = []

            def list_customerself_resource_records(self, request):
                self.requests.append(request)
                return bss_intl.ListCustomerselfResourceRecordsResponse(currency="USD", total_count=1, fee_records=[
                    bss_intl.ResFeeRecordV2(resource_id="vm-1", resource_name="web-01", region="la-north-2",
                                            region_name="LA-Mexico City2", cloud_service_type=S + "ec2",
                                            cloud_service_type_name="Elastic Cloud Server",
                                            product_spec_desc="x0.4u.8g", amount=12.5)])

        fake = Fake()
        builder = FakeBuilder(fake)
        with mock.patch.object(bss_intl.BssintlClient, "new_builder", return_value=builder), \
                mock.patch.object(bss_china.BssClient, "new_builder", side_effect=AssertionError("BSS China")):
            registros = huawei_costs.consultar_recursos_mes(FAKE_AK, FAKE_SK, "2026-09")
        self.assertEqual(builder.region.endpoints[0], INTL_ENDPOINT)
        self.assertEqual(registros, [{"resource_id": "vm-1", "resource_name": "web-01", "region": "la-north-2",
                                      "region_name": "LA-Mexico City2", "service_code": S + "ec2",
                                      "service_name": "Elastic Cloud Server", "resource_type_name": "",
                                      "spec": "x0.4u.8g", "amount": 12.5, "currency": "USD"}])
        self.assertEqual(fake.requests[0].cycle, "2026-09")

    def test_resource_errors_are_classified_and_redacted(self):
        class Denied:
            def list_customerself_resource_records(self, request):
                raise api_error(403, f"Access denied. {FAKE_AK} {FAKE_SK}", "CBC.0151", "req-9")

        with mock.patch.object(bss_intl.BssintlClient, "new_builder", return_value=FakeBuilder(Denied())):
            with self.assertRaises(huawei_costs.CostApiError) as ctx:
                huawei_costs.consultar_recursos_mes(FAKE_AK, FAKE_SK, "2026-09")
        error = ctx.exception
        self.assertEqual((error.http_status, error.error_code, error.request_id), (403, "CBC.0151", "req-9"))
        self.assertNotIn(FAKE_AK, str(error))
        self.assertNotIn(FAKE_SK, str(error))


if __name__ == "__main__":
    unittest.main()
