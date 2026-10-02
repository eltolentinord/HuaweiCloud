# coding: utf-8
"""Endpoints FastAPI con el SDK simulado (requiere httpx, ver requirements-dev.txt)."""

import io
import unittest
from types import SimpleNamespace
from unittest import mock

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.helpers import FAKE_AK, FAKE_PROJECT, FAKE_SK, REGION, FakeClient, FakeFactory, api_error
from tests.test_collectors import ECS_SERVER

from app import app
from mock_data_MOCK import MOCK_RESPONSE

RESPONSE_KEYS = {"service", "region", "project_id_masked", "total_tablas", "resumen", "errores", "avisos", "tables"}


def payload(service="ecs", **kw):
    data = {"ak": FAKE_AK, "sk": FAKE_SK, "project_id": FAKE_PROJECT, "region": REGION, "service": service}
    data.update(kw)
    return data


class TestApp(unittest.TestCase):
    def setUp(self):
        self.http = TestClient(app)

    def with_client(self, fake):
        return mock.patch("core.engine.ClientFactory", lambda creds, **kw: FakeFactory(fake))

    def test_index_renders_regions_and_15_services(self):
        response = self.http.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("ap-southeast-3", response.text)
        self.assertIn("Cloud Firewall (CFW)", response.text)

    def test_static_assets(self):
        for path in ("/static/app.js", "/static/styles.css", "/static/costs.js"):
            self.assertEqual(self.http.get(path).status_code, 200, path)

    def test_inventory_response_contract(self):
        fake = FakeClient(list_servers_details=lambda r: SimpleNamespace(servers=[ECS_SERVER], count=1))
        with self.with_client(fake):
            response = self.http.post("/api/inventory", json=payload())
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(set(body), RESPONSE_KEYS)
        self.assertEqual(body["project_id_masked"], "012345…cdef")
        self.assertEqual(body["tables"][0]["filas"][0]["Sistema operativo"], "Ubuntu 22.04 server 64bit")
        self.assertNotIn(FAKE_SK, response.text)
        self.assertNotIn(FAKE_AK, response.text)
        self.assertNotIn(FAKE_PROJECT, response.text)

    def test_401_single_service_returns_502_with_safe_detail(self):
        def handler(request):
            raise api_error(401, f"bad {FAKE_SK}", "APIGW.0301")
        with self.with_client(FakeClient(list_servers_details=handler)):
            response = self.http.post("/api/inventory", json=payload())
        self.assertEqual(response.status_code, 502)
        detail = response.json()["detail"]
        self.assertTrue(detail["mensaje"].startswith("Credenciales no válidas"))
        self.assertEqual(detail["error_code"], "APIGW.0301")
        self.assertNotIn(FAKE_SK, response.text)

    def test_403_returns_warning(self):
        def handler(request):
            raise api_error(403, "Forbidden")
        with self.with_client(FakeClient(list_servers_details=handler)):
            response = self.http.post("/api/inventory", json=payload())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["avisos"][0]["tipo"], "aviso")

    def test_invalid_service_and_blank_credentials(self):
        self.assertEqual(self.http.post("/api/inventory", json=payload("nope")).status_code, 400)
        self.assertEqual(self.http.post("/api/inventory", json=payload(ak="  ")).status_code, 400)
        self.assertEqual(self.http.post("/api/inventory", json=payload(sk="")).status_code, 422)

    def test_export_mock_tables(self):
        body = {"region": REGION, "service": "ecs", "tables": MOCK_RESPONSE["tables"],
                "avisos": [{"servicio": "HSS", "mensaje": "=evil"}]}
        response = self.http.post("/api/export", json=body)
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="inventario_ecs_ap-southeast-3.xlsx"', response.headers["content-disposition"])
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["ECS", "Avisos"])

    def test_export_sanitizes_filename(self):
        body = {"region": "x\"; evil", "service": "ecs", "tables": MOCK_RESPONSE["tables"]}
        response = self.http.post("/api/export", json=body)
        self.assertIn('filename="inventario_ecs_x___evil.xlsx"', response.headers["content-disposition"])

    def test_export_requires_tables(self):
        response = self.http.post("/api/export", json={"region": REGION, "service": "ecs", "tables": []})
        self.assertEqual(response.status_code, 400)


class TestCostRoutesStillWork(unittest.TestCase):
    def setUp(self):
        self.http = TestClient(app)

    def test_costs_page(self):
        self.assertEqual(self.http.get("/costos").status_code, 200)

    def test_compare_validation_without_network(self):
        response = self.http.post("/api/costs/compare",
                                  json={"ak": "", "sk": "", "month_a": "", "month_b": ""})
        self.assertTrue(response.json()["error"])

    def test_compare_errors_hide_secrets(self):
        def boom(*args, **kwargs):
            raise RuntimeError(f"failed with {FAKE_SK}")
        with mock.patch("routers.costs.consultar_mes", boom):
            response = self.http.post("/api/costs/compare", json={
                "ak": FAKE_AK, "sk": FAKE_SK, "month_a": "2026-01", "month_b": "2026-02"})
        self.assertTrue(response.json()["error"])
        self.assertNotIn(FAKE_SK, response.text)


if __name__ == "__main__":
    unittest.main()
