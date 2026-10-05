# coding: utf-8
"""Auditoría de servidores por SSH: CRUD con contraseña cifrada, ejecución del script del
usuario (sudo por stdin, huella SSH, limpieza), lotes y reportes. Sin red: SSH simulado."""

import io
import json
import unittest
import uuid
from datetime import datetime
from types import SimpleNamespace
from unittest import mock

from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_admin_api import AdminApiTestCase

from core.authz import Principal
from db.models import AuditEvent, Server, ServerAuditRun
from db.session import get_db_session_factory, session_scope
from routers import servers_api
from routers.security import get_principal
from server_audit import report as reports
from server_audit import service
from server_audit.runner import ExecResult, SshError, fingerprint_of
from server_audit.servers import decrypt_password

PASSWORD = "Sup3r-Secreta!#pass"
NEW_PASSWORD = "Otra-Clave-2026$"
FP = "SHA256:abcdefFAKEfingerprint0123456789"


def audit_json(hostname="web-01", score=72, fail=True):
    return {
        "hostname": hostname, "generated_at": "2026-10-05 12:00:00", "os": "Ubuntu 22.04.4 LTS",
        "overall_score": score, "overall_points": {"got": 108, "max": 150},
        "hardening": {"got": 60, "max": 100, "pct": 60}, "updates": {"got": 48, "max": 50, "pct": 96},
        "checks": [
            {"id": "hd_firewall", "category": "hardening", "description": "Firewall activo", "status": "OK",
             "points": 15, "max": 15, "recommendation": ""},
            {"id": "hd_ssh_passwordauth", "category": "hardening", "description": "SSH: password deshabilitado",
             "status": "FAIL" if fail else "OK", "points": 0 if fail else 15, "max": 15,
             "recommendation": "Configura 'PasswordAuthentication no'." if fail else ""},
        ],
        "sites": [{"name": "www.example.org", "source": "nginx"}],
    }


class FakeSession:
    """Sesión SSH simulada: registra comandos/stdin y devuelve los archivos que 'genera' el script."""

    def __init__(self, *, fingerprint=FP, data=None, html="<html><body>reporte</body></html>",
                 sudo_ok=True, script_ok=True):
        self.fingerprint = fingerprint
        self.commands, self.stdins, self.uploads, self.closed = [], [], {}, False
        self.files = {}
        self.data, self.html, self.sudo_ok, self.script_ok = data or audit_json(), html, sudo_ok, script_ok

    def run(self, command, *, stdin=None, timeout=60):
        self.commands.append(command)
        self.stdins.append(stdin)
        if "server-audit.sh" in command:
            if "sudo" in command and not self.sudo_ok:
                return ExecResult(1, "", "Sorry, try again.\nsudo: 3 incorrect password attempts")
            if not self.script_ok:
                return ExecResult(2, "", f"línea 12: error con {PASSWORD}")
            remote = command.split(" -j ")[1].split()[0].rsplit("/", 1)[0]
            self.files[f"{remote}/r.json"] = json.dumps(self.data).encode()
            self.files[f"{remote}/r.html"] = self.html.encode()
            return ExecResult(0, "ok", "")
        if command.startswith("sudo") and "true" in command:
            return ExecResult(0 if self.sudo_ok else 1, "", "")
        if command.startswith("command -v bash"):
            return ExecResult(0, "ok\n", "")
        return ExecResult(0, "", "")

    def put(self, data, remote_path, *, mode=0o700):
        self.uploads[remote_path] = (data, mode)

    def get(self, remote_path):
        if remote_path not in self.files:
            raise IOError("no existe")
        return self.files[remote_path]

    def close(self):
        self.closed = True


class ServerApiTestCase(AdminApiTestCase):
    def setUp(self):
        super().setUp()
        self.app.dependency_overrides[get_db_session_factory] = lambda: self.factory
        patcher = mock.patch.object(service, "MAX_PARALLEL", 1)   # SQLite en memoria: sin hilos en tests
        patcher.start()
        self.addCleanup(patcher.stop)
        self.sessions = []
        self.behaviour = {}
        self.connects = []

        def connect(**kwargs):
            self.connects.append(kwargs)
            if self.behaviour.get("raise"):
                raise self.behaviour["raise"]
            expected = kwargs.get("expected_fingerprint")
            fp = self.behaviour.get("fingerprint", FP)
            if expected and expected != fp:
                raise SshError("fingerprint", "La huella del servidor cambió: posible suplantación.")
            session = FakeSession(fingerprint=fp, **self.behaviour.get("session", {}))
            self.sessions.append(session)
            return session

        servers_api.CONNECTOR["connect"] = connect
        self.cid = self.new_client("Itla")["id"]
        self.base = f"/api/clients/{self.cid}/servers"

    def tearDown(self):
        servers_api.CONNECTOR["connect"] = None
        for response in self.responses:
            self.assertNotIn(PASSWORD, response.text)
            self.assertNotIn(NEW_PASSWORD, response.text)
        logs = self.log_stream.getvalue()
        self.assertNotIn(PASSWORD, logs)
        self.assertNotIn(NEW_PASSWORD, logs)
        super().tearDown()

    def new_server(self, name="web-01", host="10.0.0.5", **extra):
        body = {"name": name, "host": host, "port": 22, "username": "auditor", "password": PASSWORD, **extra}
        response = self.call("POST", self.base, json=body)
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def as_role(self, role):
        self.app.dependency_overrides[get_principal] = lambda: Principal(
            subject=f"u-{role}", kind="user", client_roles={self.cid: role})


class TestServerCrud(ServerApiTestCase):
    def test_password_is_encrypted_and_never_returned(self):
        server = self.new_server()
        self.assertTrue(server["has_password"])
        self.assertNotIn("password", {k for k in server if k != "has_password"})
        with session_scope(self.factory) as db:
            row = db.get(Server, uuid.UUID(server["id"]))
            self.assertNotIn(PASSWORD.encode(), bytes(row.password_ciphertext))
            self.assertEqual(decrypt_password(row, self.keyring), PASSWORD)
            events = list(db.scalars(select(AuditEvent).where(AuditEvent.action == "server.create")))
            self.assertEqual(len(events), 1)
            self.assertNotIn(PASSWORD, json.dumps(events[0].details, default=str))

    def test_edit_keeps_password_when_empty_and_replaces_it(self):
        server = self.new_server()
        body = {"name": "web-01b", "host": "web01.example.org", "port": 2222, "username": "auditor", "password": ""}
        updated = self.call("PUT", f"{self.base}/{server['id']}", json=body).json()
        self.assertEqual((updated["name"], updated["host"], updated["port"]), ("web-01b", "web01.example.org", 2222))
        with session_scope(self.factory) as db:
            self.assertEqual(decrypt_password(db.get(Server, uuid.UUID(server["id"])), self.keyring), PASSWORD)
        self.assertEqual(self.call("PUT", f"{self.base}/{server['id']}/password",
                                   json={"password": NEW_PASSWORD}).status_code, 200)
        with session_scope(self.factory) as db:
            self.assertEqual(decrypt_password(db.get(Server, uuid.UUID(server["id"])), self.keyring), NEW_PASSWORD)

    def test_validation_and_duplicates(self):
        self.new_server()
        bad = [{"host": "http://x"}, {"host": "a b"}, {"username": "Root!"}, {"port": 0}, {"password": ""},
               {"name": ""}]
        for extra in bad:
            with self.subTest(extra=extra):
                body = {"name": "x", "host": "10.0.0.9", "port": 22, "username": "auditor", "password": PASSWORD, **extra}
                self.assertIn(self.call("POST", self.base, json=body).status_code, (409, 422))
        self.assertEqual(self.call("POST", self.base, json={"name": "web-01", "host": "10.0.0.6", "port": 22,
                                                            "username": "auditor", "password": PASSWORD}).status_code, 409)

    def test_isolation_and_roles(self):
        server = self.new_server()
        other = self.new_client("Otra")["id"]
        self.assertEqual(self.call("GET", f"/api/clients/{other}/servers").json()["servers"], [])
        self.assertEqual(self.call("PUT", f"/api/clients/{other}/servers/{server['id']}/password",
                                   json={"password": "x"}).status_code, 404)
        self.as_role("viewer")
        self.assertEqual(self.call("GET", self.base).status_code, 200)
        self.assertEqual(self.call("POST", f"{self.base}/audits", json={}).status_code, 403)
        self.assertEqual(self.call("POST", self.base, json={}).status_code, 403)
        self.as_role("operator")
        self.assertEqual(self.call("POST", f"{self.base}/audits", json={}).status_code, 202)
        self.assertEqual(self.call("DELETE", f"{self.base}/{server['id']}").status_code, 403)


class TestRunAudit(ServerApiTestCase):
    def test_audit_runs_script_with_sudo_stdin_and_cleans_up(self):
        server = self.new_server()
        response = self.call("POST", f"{self.base}/audits", json={"server_ids": [server["id"]]})
        self.assertEqual(response.status_code, 202, response.text)
        session = self.sessions[0]
        script_cmd = next(c for c in session.commands if "server-audit.sh" in c)
        self.assertTrue(script_cmd.startswith("sudo -k -S -p '' bash /tmp/hci-audit-"))
        self.assertNotIn(PASSWORD, script_cmd)                                   # nunca en la línea de comandos
        self.assertEqual(session.stdins[session.commands.index(script_cmd)], PASSWORD + "\n")
        uploaded = next(iter(session.uploads.values()))
        self.assertEqual((uploaded[0], uploaded[1]), (service.script_bytes(), 0o700))
        self.assertTrue(any(c.startswith("rm -rf -- /tmp/hci-audit-") for c in session.commands))
        self.assertTrue(session.closed)
        self.assertIsNone(self.connects[0]["expected_fingerprint"])
        listed = self.call("GET", self.base).json()["servers"][0]
        self.assertEqual(listed["host_fingerprint"], FP)                         # TOFU: primera huella guardada
        last = listed["last_audit"]
        self.assertEqual((last["status"], last["overall_score"], last["hardening_pct"], last["updates_pct"]),
                         ("succeeded", 72, 60, 96))
        detail = self.call("GET", f"{self.base}/audits/{last['id']}").json()
        self.assertEqual(detail["result"]["sites"][0]["name"], "www.example.org")
        page = self.call("GET", f"{self.base}/audits/{last['id']}/report.html")
        self.assertIn("reporte", page.text)
        self.assertIn("default-src 'none'", page.headers["content-security-policy"])

    def test_changed_fingerprint_is_rejected(self):
        server = self.new_server()
        self.call("POST", f"{self.base}/audits", json={})
        self.behaviour["fingerprint"] = "SHA256:otraHuella"
        self.call("POST", f"{self.base}/audits", json={})
        self.assertEqual(self.connects[1]["expected_fingerprint"], FP)
        last = self.call("GET", self.base).json()["servers"][0]["last_audit"]
        self.assertEqual((last["status"], last["error_kind"]), ("failed", "fingerprint"))
        self.assertEqual(self.call("POST", f"{self.base}/{server['id']}/fingerprint/reset").status_code, 200)
        self.call("POST", f"{self.base}/audits", json={})
        self.assertEqual(self.call("GET", self.base).json()["servers"][0]["host_fingerprint"], "SHA256:otraHuella")

    def test_errors_are_classified_and_redacted(self):
        self.new_server()
        cases = [({"raise": SshError("auth", "Usuario o contraseña SSH incorrectos.")}, "auth"),
                 ({"session": {"sudo_ok": False}}, "sudo"),
                 ({"session": {"script_ok": False}}, "script"),
                 ({"session": {"data": {"hostname": "x"}}}, "result")]
        for behaviour, kind in cases:
            with self.subTest(kind=kind):
                self.behaviour = behaviour
                self.call("POST", f"{self.base}/audits", json={})
                last = self.call("GET", self.base).json()["servers"][0]["last_audit"]
                self.assertEqual((last["status"], last["error_kind"]), ("failed", kind))
                self.assertNotIn(PASSWORD, last["error_message"] or "")

    def test_root_runs_without_sudo(self):
        self.new_server(username="root")
        self.call("POST", f"{self.base}/audits", json={})
        script_cmd = next(c for c in self.sessions[0].commands if "server-audit.sh" in c)
        self.assertTrue(script_cmd.startswith("bash /tmp/hci-audit-"))
        self.assertIsNone(self.sessions[0].stdins[self.sessions[0].commands.index(script_cmd)])

    def test_batch_continues_after_a_failure_and_reports(self):
        self.new_server("web-01", "10.0.0.5")
        self.new_server("db-01", "10.0.0.6")
        original = servers_api.CONNECTOR["connect"]

        def flaky(**kwargs):
            if kwargs["host"] == "10.0.0.6":
                raise SshError("network", "No se pudo conectar a 10.0.0.6:22.")
            return original(**kwargs)
        servers_api.CONNECTOR["connect"] = flaky
        batch = self.call("POST", f"{self.base}/audits", json={}).json()
        runs = self.call("GET", f"{self.base}/audits", params={"batch_id": batch["batch_id"]}).json()
        self.assertEqual(sorted(r["status"] for r in runs["runs"]), ["failed", "succeeded"])
        self.assertEqual(runs["pending"], 0)
        html = self.call("GET", f"{self.base}/audits/report", params={"format": "html", "batch_id": batch["batch_id"]})
        self.assertIn("web-01", html.text)
        self.assertIn("Firewall activo", html.text)
        pdf = self.call("GET", f"{self.base}/audits/report", params={"format": "pdf"})
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        xlsx = self.call("GET", f"{self.base}/audits/report", params={"format": "xlsx"})
        wb = load_workbook(io.BytesIO(xlsx.content))
        self.assertEqual(wb.sheetnames, ["Resumen", "Controles", "Recomendaciones", "Sitios"])

    def test_test_connection_saves_first_fingerprint(self):
        server = self.new_server()
        result = self.call("POST", f"{self.base}/{server['id']}/test").json()
        self.assertEqual((result["ok"], result["fingerprint"], result["sudo"]), (True, FP, True))
        self.assertEqual(self.call("GET", self.base).json()["servers"][0]["host_fingerprint"], FP)
        self.behaviour["raise"] = SshError("network", "No se pudo conectar.")
        failed = self.call("POST", f"{self.base}/{server['id']}/test").json()
        self.assertEqual((failed["ok"], failed["error_kind"]), (False, "network"))

    def test_report_requires_successful_audits(self):
        self.assertEqual(self.call("GET", f"{self.base}/audits/report").status_code, 422)


class TestReports(unittest.TestCase):
    def test_comparative_sorted_and_escaped(self):
        results = [{"server_name": "bajo", "data": audit_json("low", 40)},
                   {"server_name": "<b>alto</b>", "data": audit_json("high", 90, fail=False)}]
        html = reports.comparative_html(results, generated_at=datetime(2026, 10, 5))
        self.assertLess(html.index("&lt;b&gt;alto"), html.index("bajo"))
        self.assertNotIn("<b>alto</b>", html)

    def test_fingerprint_format(self):
        key = SimpleNamespace(asbytes=lambda: b"clave-publica")
        self.assertTrue(fingerprint_of(key).startswith("SHA256:"))
        self.assertNotIn("=", fingerprint_of(key))

    def test_real_script_is_bundled_unchanged(self):
        self.assertTrue(service.script_bytes().startswith(b"#!/usr/bin/env bash"))
        self.assertNotIn(b"\r\n", service.script_bytes())
        self.assertEqual(len(service.script_sha256()), 64)


class TestServersPage(unittest.TestCase):
    def test_page_and_sidebar(self):
        from fastapi.testclient import TestClient
        from app import app
        http = TestClient(app)
        page = http.get("/servidores").text
        for marker in ('id="serverTable"', 'id="serverForm"', 'type="password"', 'autocomplete="new-password"',
                       'src="/static/servers.js?v='):
            self.assertIn(marker, page)
        self.assertIn('href="/servidores"', http.get("/clientes").text)


if __name__ == "__main__":
    unittest.main()
