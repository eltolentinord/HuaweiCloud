# coding: utf-8
"""Exportaciones del inventario persistido (Excel/CSV): contenido, filtros, diferencias,
formula injection, aislamiento y ausencia de secretos."""

import csv
import io
import unittest
import zipfile

from openpyxl import load_workbook

from tests.helpers import FAKE_AK, FAKE_SK
from tests.scan_helpers import server
from tests.test_inventory_api import InventoryApiTestCase

from core.authz import Principal
from exports.inventory_export import csv_safe, to_csv
from routers.security import get_principal

EVIL = ["=HYPERLINK(\"http://x\")", "+1+1", "-2+3", "@SUM(A1)", "\t=cmd", "\r=cmd"]


class TestCsvSafety(unittest.TestCase):
    def test_csv_prefixes_dangerous_values(self):
        for value in EVIL:
            self.assertTrue(csv_safe(value).startswith("'"), repr(value))
        for value in ("web", "10.0.0.1", "", None, 5):
            self.assertFalse(csv_safe(value).startswith("'"))

    def test_csv_file_is_utf8_bom_and_quoted(self):
        data = to_csv(["a", "b"], [{"a": "ñandú", "b": "=1+1"}]).getvalue()
        self.assertTrue(data.startswith("﻿".encode("utf-8")))
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(rows, [["a", "b"], ["ñandú", "'=1+1"]])


class ExportApiTestCase(InventoryApiTestCase):
    def export(self, kind, **params):
        response = self.call("GET", f"{self.base}/exports/{kind}", params=params)
        self.assertEqual(response.status_code, 200, response.text[:300])
        return response

    @staticmethod
    def xlsx(response):
        return load_workbook(io.BytesIO(response.content))

    @staticmethod
    def csv_rows(response):
        return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))


class TestInventoryExports(ExportApiTestCase):
    def test_full_inventory_xlsx(self):
        response = self.export("inventory")
        self.assertIn("attachment;", response.headers["content-disposition"])
        self.assertTrue(response.headers["content-disposition"].endswith('.xlsx"'))
        ws = self.xlsx(response)["Recursos"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[:4], ["Servicio", "Tipo", "ID proveedor", "Nombre"])
        self.assertEqual(ws.max_row - 1, 5)  # 4 ECS + 1 bucket
        self.assertEqual(ws.freeze_panes, "A2")

    def test_filtered_csv(self):
        rows = self.csv_rows(self.export("inventory", format="csv", service="ecs", search="web"))
        self.assertEqual({r["ID proveedor"] for r in rows}, {"srv-1", "srv-3"})
        self.assertEqual({r["Project ID"] for r in rows}, {"p-sg"})

    def test_deleted_only_export(self):
        self.world.servers["p-sg"] = []
        self.scan()
        rows = self.csv_rows(self.export("inventory", format="csv", only_deleted="true"))
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(r["Eliminado"] for r in rows))

    def test_summary(self):
        rows = self.csv_rows(self.export("summary", format="csv"))
        values = {(r["Grupo"], r["Clave"]): r["Valor"] for r in rows}
        self.assertEqual(values[("Total", "Recursos activos")], "5")
        self.assertEqual(values[("Servicio", "ecs")], "4")
        self.assertEqual(values[("Último escaneo", "Estado")], "completed")
        self.assertIn("Resumen", self.xlsx(self.export("summary")).sheetnames)

    def test_compare_exports(self):
        self.world.servers["p-sg"] = [server("srv-1", "renamed"), server("srv-9", "new")]
        second = self.scan()
        wb = self.xlsx(self.export("compare", from_scan=self.first, to_scan=second))
        self.assertEqual(wb.sheetnames, ["Resumen", "Agregados", "Eliminados", "Modificados", "Transitorios"])
        self.assertEqual(wb["Agregados"]["C2"].value, "srv-9")
        modified = [[c.value for c in row] for row in wb["Modificados"].iter_rows(min_row=2)]
        self.assertTrue(any("name" in row and "renamed" in row for row in modified))
        rows = self.csv_rows(self.export("compare", format="csv", from_scan=self.first, to_scan=second))
        self.assertEqual({r["Categoría"] for r in rows}, {"Agregados", "Eliminados", "Modificados"})


class TestExportSafety(ExportApiTestCase):
    def test_formula_injection_in_both_formats(self):
        self.world.servers["p-mx"] = [server("srv-evil", value) for value in EVIL[:1]] + \
            [server(f"srv-e{i}", value) for i, value in enumerate(EVIL[1:4])]
        self.scan()
        xlsx = self.export("inventory", service="ecs")
        with zipfile.ZipFile(io.BytesIO(xlsx.content)) as archive:
            sheet = "".join(archive.read(n).decode() for n in archive.namelist() if n.startswith("xl/worksheets/"))
        self.assertNotIn("<f>", sheet)
        names = {r["Nombre"] for r in self.csv_rows(self.export("inventory", format="csv", region="la-north-2"))}
        self.assertTrue(names)
        self.assertTrue(all(n.startswith("'") for n in names), names)

    def test_no_raw_and_no_secrets(self):
        self.world.servers["p-sg"] = [server("srv-1", "web", metadata={"note": FAKE_SK},
                                             **{"OS-EXT-SRV-ATTR:user_data": "IyEvYmluL2Jhc2g="})]
        self.scan()
        for fmt in ("xlsx", "csv"):
            response = self.export("inventory", format=fmt)
            payload = response.content if fmt == "csv" else b"".join(
                zipfile.ZipFile(io.BytesIO(response.content)).read(n)
                for n in zipfile.ZipFile(io.BytesIO(response.content)).namelist())
            for secret in (FAKE_SK.encode(), FAKE_AK.encode(), b"IyEvYmluL2Jhc2g="):
                self.assertNotIn(secret, payload, fmt)

    def test_isolation_and_permissions(self):
        other = f"/api/clients/{self.other_cid}/accounts/{self.aid}/exports/inventory"
        self.assertEqual(self.call("GET", other).status_code, 404)
        self.app.dependency_overrides[get_principal] = lambda: Principal(
            subject="v", kind="user", client_roles={str(self.cid): "viewer"})
        self.assertEqual(self.call("GET", f"{self.base}/exports/summary").status_code, 200)  # viewer exporta
        self.assertEqual(self.call("GET", f"/api/clients/{self.other_cid}/accounts/{self.other_aid}/exports/summary"
                                   ).status_code, 404)

    def test_invalid_parameters(self):
        self.assertEqual(self.call("GET", f"{self.base}/exports/inventory", params={"format": "pdf"}).status_code, 422)
        self.assertEqual(self.call("GET", f"{self.base}/exports/inventory", params={"sort": "nope"}).status_code, 422)
